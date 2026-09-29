"""
AEOS First IN / Last OUT report - standalone Excel generator.
Fallback for when the Jasper report inside AEOS isn't available.

Usage:
    python generate_first_in_last_out.py 2026-06-24 2026-06-25
    python generate_first_in_last_out.py 2026-06-24 2026-06-25 --out report.xlsx
"""

import argparse
import sys
from datetime import datetime, timedelta

import pyodbc
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SERVER = "localhost"
DATABASE = "aeosdb"

QUERY = """
SELECT
    c.objectid                          AS carrier_id,
    LTRIM(RTRIM(c.lastname))            AS last_name,
    LTRIM(RTRIM(c.initials))            AS initials,
    c.personnelnr                       AS file_number,
    dep.name                            AS department,
    badges.card_number                  AS card_number,
    CAST(el.timestamp AS DATE)          AS activity_date,
    MIN(CASE WHEN el.intvalue = 1 THEN el.timestamp END) AS first_in,
    MAX(CASE WHEN el.intvalue = 2 THEN el.timestamp END) AS last_out,
    DATEDIFF(MINUTE,
        MIN(CASE WHEN el.intvalue = 1 THEN el.timestamp END),
        MAX(CASE WHEN el.intvalue = 2 THEN el.timestamp END)) AS duration_minutes,
    COUNT(el.objectid)                  AS event_count
FROM        eventlog         el
JOIN        eventtype        et  ON et.objectid  = el.eventtype
                                AND et.eventcategory = 1
JOIN        carrier          c   ON c.objectid   = el.carrierid
                                AND c.carriertype = 1
                                AND c.removaldate IS NULL
LEFT JOIN   department       dep ON dep.objectid  = c.departmentobjectid
OUTER APPLY (
    SELECT STRING_AGG(x.badgenumber, ', ') AS card_number
    FROM (
        SELECT DISTINCT tk.badgenumber
        FROM   tokenassignment ta
        JOIN   token           tk ON tk.objectid = ta.identifierobjectid
        WHERE  ta.carrierobjectid = c.objectid AND ta.withdrawn = 0
    ) x
) badges
WHERE
    el.timestamp >= ?
    AND el.timestamp <= ?
GROUP BY
    c.objectid, c.lastname, c.initials, c.personnelnr,
    dep.name, badges.card_number, CAST(el.timestamp AS DATE)
ORDER BY c.lastname, c.initials, CAST(el.timestamp AS DATE)
"""

COLUMNS = [
    ("File No", "file_number", 12),
    ("Last Name", "last_name", 20),
    ("First Name", "initials", 18),
    ("Department", "department", 30),
    ("Card No", "card_number", 16),
    ("Date", "activity_date", 12),
    ("First In", "first_in", 12),
    ("Last Out", "last_out", 12),
    ("Duration (HH:MM)", "duration_minutes", 16),
    ("Events", "event_count", 8),
]


def fetch_rows(date_from, date_to):
    conn_str = (
        f"DRIVER={{ODBC Driver 18 for SQL Server}};"
        f"SERVER={SERVER};DATABASE={DATABASE};"
        f"Trusted_Connection=yes;Encrypt=no;"
    )
    try:
        conn = pyodbc.connect(conn_str)
    except pyodbc.Error:
        # fall back to the older driver name if 18 isn't installed
        conn_str = conn_str.replace("ODBC Driver 18 for SQL Server", "SQL Server")
        conn = pyodbc.connect(conn_str)

    cursor = conn.cursor()
    cursor.execute(QUERY, date_from, date_to)
    columns = [c[0] for c in cursor.description]
    rows = [dict(zip(columns, row)) for row in cursor.fetchall()]
    conn.close()
    return rows


def format_duration(minutes):
    if minutes is None or minutes < 0:
        return "—"
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def build_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "First In Last Out"

    header_fill = PatternFill(start_color="1A3A5C", end_color="1A3A5C", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, (label, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (_, key, _) in enumerate(COLUMNS, start=1):
            value = row.get(key)
            if key == "duration_minutes":
                value = format_duration(value)
            elif key in ("first_in", "last_out") and value is not None:
                value = value.strftime("%H:%M:%S")
            ws.cell(row=row_idx, column=col_idx, value=value)

    ws.freeze_panes = "A2"
    return wb


def main():
    parser = argparse.ArgumentParser(description="Generate AEOS First In/Last Out Excel report")
    parser.add_argument("date_from", help="Start date, YYYY-MM-DD")
    parser.add_argument("date_to", help="End date (inclusive), YYYY-MM-DD")
    parser.add_argument("--out", default=None, help="Output .xlsx path")
    args = parser.parse_args()

    date_from = datetime.strptime(args.date_from, "%Y-%m-%d")
    date_to = datetime.strptime(args.date_to, "%Y-%m-%d") + timedelta(days=1) - timedelta(seconds=1)

    print(f"Connecting to {SERVER}/{DATABASE} (Windows auth)...")
    rows = fetch_rows(date_from, date_to)
    print(f"Fetched {len(rows)} rows.")

    wb = build_workbook(rows)
    out_path = args.out or f"FirstInLastOut_{args.date_from}_to_{args.date_to}.xlsx"
    wb.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    sys.exit(main())
